from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / "课程报告.docx"
OUTPUT = ROOT / "课程报告-已补全6.6和6.7.docx"
DOWNLOADS = Path.home() / "Downloads"

FIG_REORDER = DOWNLOADS / "累计确认与失序重组.png"
FIG_RTT = DOWNLOADS / "正常 RTT采样.png"
FIG_FAST = DOWNLOADS / "快速重传.png"


def find_paragraph(doc, text):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


def move_before(anchor, element):
    anchor._p.addprevious(element)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_paragraph(doc, anchor, text="", style="Normal", first_indent=False):
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)
    if first_indent:
        paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.space_after = Pt(6)
    move_before(anchor, paragraph._p)
    return paragraph


def add_labeled_paragraph(doc, anchor, label, text):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.add_run(label).bold = True
    paragraph.add_run(text)
    paragraph.paragraph_format.space_after = Pt(6)
    move_before(anchor, paragraph._p)
    return paragraph


def add_heading(doc, anchor, text, level=3):
    paragraph = add_paragraph(doc, anchor, text, style=f"Heading {level}")
    set_keep_with_next(paragraph)
    return paragraph


def add_picture(doc, anchor, path, width, caption):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    paragraph.paragraph_format.space_after = Pt(3)
    set_keep_with_next(paragraph)
    move_before(anchor, paragraph._p)

    cap = add_paragraph(doc, anchor, caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_result_table(doc, anchor, rows):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["测试项", "条件与流程", "预期结果", "实际结果与判定"]
    for index, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], value, bold=True)
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            set_cell_text(cells[index], value)
    move_before(anchor, table._tbl)
    spacer = add_paragraph(doc, anchor)
    spacer.paragraph_format.space_after = Pt(3)
    return table


def clear_section_body(doc, heading_text, next_heading_text):
    heading = find_paragraph(doc, heading_text)
    next_heading = find_paragraph(doc, next_heading_text)
    node = heading._p.getnext()
    while node is not None and node is not next_heading._p:
        following = node.getnext()
        node.getparent().remove(node)
        node = following


def update_summary_table(doc):
    table = doc.tables[11]
    values = {
        1: [
            "累计确认/失序重组",
            "以 rcv_nxt 表示下一期待字节；ACK 累计确认连续前缀；失序段按序号插入队列，缺口补齐后再按序交付，并抑制重复数据。",
            "src/tju_tcp.c：insert_out_of_order_locked()、drain_ordered_locked()、tju_handle_packet()",
            "RDT-REL-01：100 Mbps，100 ms 时延，80 ms 抖动，5% 丢包，25% 乱序",
            "图表7；机制证据通过，但 cmp 返回1，端到端字节一致性未通过。",
        ],
        2: [
            "快速/超时重传",
            "3次重复ACK重传首个未确认段；RTO到期重传并指数退避，重传样本不参与RTT估计。",
            "src/tju_tcp.c：process_ack()、timer_main()、retransmit_segment()",
            "RDT-REL-02：100 Mbps，20 ms 时延，10% 丢包；统计重复ACK与重复SEND序号",
            "图表8；观察到重复ACK及同序号重发。截图未保留时序，快速重传时机证据不完整。",
        ],
        3: [
            "RTT/RTO",
            "按 RFC 6298 更新 SRTT、RTTVAR 与 RTO；RTO限制在200~4000 ms；采用 Karn 算法排除重传样本。",
            "src/tju_tcp.c：update_rtt_locked()、process_ack()、timer_main()",
            "RDT-REL-03：基准链路采集 client.event.trace 中 RTTS 事件",
            "图表9；数值与平滑公式及最小RTO钳位一致，通过。",
        ],
        4: [
            "流量控制/零窗口",
            "接收端通告可用缓存；发送端受 snd_una+peer_rwnd 右边界限制；零窗口时停止新数据并周期探测，窗口恢复后续传。",
            "src/tju_tcp.c：advertised_window_locked()、flush_send_queue()、send_window_probe()、tju_recv()",
            "RDT-FLOW-01正常窗口；RDT-FLOW-02暂停读取压满缓存后恢复",
            "6.7节；静态路径核验完成，现有运行未出现 rwnd=0，零窗口动态测试未通过验收。",
        ],
    }
    for row_index, row_values in values.items():
        for col_index, value in enumerate(row_values):
            set_cell_text(table.rows[row_index].cells[col_index], value)


def build_section_66(doc, anchor):
    add_paragraph(
        doc,
        anchor,
        "本节使用 test/rdt_client、test/rdt_server 以及 client.event.trace、server.event.trace 验证累计确认、失序重组、重传和 RTT/RTO。数据段最大载荷为 1375 B。除特别说明外，测试前清空 trace，先启动服务端，再设置两端链路参数并运行客户端；测试后用 grep/sed/uniq 检查 ACK 与重发序号，用 DELV 事件检查交付顺序，并以 cmp 检查发送文件和接收文件的逐字节一致性。",
        first_indent=True,
    )

    add_result_table(
        doc,
        anchor,
        [
            (
                "RDT-REL-01\n累计确认与失序重组",
                "100 Mbps；时延100 ms；抖动80 ms；丢包5%；乱序25%。统计重复ACK/重复SEND，检查DELV并执行cmp。",
                "缺口期间重复返回同一累计ACK；缺失段到达后连续交付；无重复交付；cmp返回0。",
                "重复ACK、重复发送及顺序DELV均出现；cmp返回1。部分通过，文件完整性失败。",
            ),
            (
                "RDT-REL-02\n快速/超时重传",
                "100 Mbps；时延20 ms；丢包10%。统计ACK出现次数与相同seq的SEND次数。",
                "3次重复ACK触发快速重传；未形成3次重复ACK时由RTO恢复。",
                "同一ACK出现12~17次，同一seq发送5~7次。确认发生丢包恢复，但截图不足以区分快速重传与RTO时机。",
            ),
            (
                "RDT-REL-03\nRTT/RTO",
                "基准链路运行RDT，提取client.event.trace前20条RTTS记录。",
                "SRTT/RTTVAR按RFC 6298平滑更新；RTO受200~4000 ms边界约束；重传样本不采样。",
                "首个有效样本40.782 ms，后续估计值与公式一致，RTO钳位为200 ms。通过。",
            ),
        ],
    )

    add_heading(doc, anchor, "6.6.1 累计确认与失序重组", level=3)
    add_labeled_paragraph(
        doc,
        anchor,
        "测试条件与流程：",
        "将两端链路设置为 100 Mbps、100 ms 时延、80 ms 抖动、5% 丢包和 25% 乱序，清空 client.event.trace 后运行 RDT 客户端。随后统计纯 ACK（flag=4、length=0）的确认号出现次数、带数据 SEND 的序号出现次数，抽查服务端 DELV 记录，最后执行 cmp test/rdt_send_file.txt test/rdt_recv_file.txt。",
    )
    add_labeled_paragraph(
        doc,
        anchor,
        "预期结果：",
        "接收端在缺口未补齐时重复发送相同的累计 ACK；发送端重发缺失段；失序段只能在前序数据到达后按序交付，每个字节只交付一次；最终 cmp 返回 0。",
    )
    add_labeled_paragraph(
        doc,
        anchor,
        "实际结果与分析：",
        "日志中多个 ACK 值重复 11~17 次，多个数据序号重复发送 7~9 次，说明乱序/丢包条件下接收端维持累计确认且发送端执行了重传。服务端 DELV 从 seq=1001 开始，通常按 1375 B 推进；在 seq=20626、30226 等位置出现 375 B 边界段后仍继续顺序推进，截图范围内未见重复交付，能够支持失序缓存补洞后按序交付的判断。但是最终 cmp 明确返回 1，并报告两个文件从第 1 字节即不同，因此端到端可靠传输不能判为通过。当前结论为“机制路径部分通过、文件完整性失败”。后续应在同一轮测试前删除旧接收文件，并保存 wc -c、sha256sum 与 cmp -l | head 的结果，以定位首个错误字节及是否存在旧文件、偏移或返回长度处理问题。",
    )
    add_picture(doc, anchor, FIG_REORDER, 5.05, "图表 7  累计确认、重复发送、按序交付及文件一致性检查")

    add_heading(doc, anchor, "6.6.2 快速重传与超时恢复", level=3)
    add_labeled_paragraph(
        doc,
        anchor,
        "测试条件与流程：",
        "将链路设置为 100 Mbps、20 ms 时延和 10% 丢包，运行 RDT 客户端；分别统计相同累计 ACK 和相同数据 seq 的出现次数。实现中 process_ack() 在同一 snd_una 累计收到 3 次重复 ACK 后重传首个未确认段；timer_main() 在 RTO 到期后同样重传该段并把 RTO 加倍，最大不超过 4000 ms。",
    )
    add_labeled_paragraph(
        doc,
        anchor,
        "预期结果：",
        "形成 3 次重复 ACK 时应在 RTO 到期前重发缺失段；若重复 ACK 不足，则由 RTO 超时恢复。重传不应推进 snd_una，也不应作为新的 RTT 样本。",
    )
    add_labeled_paragraph(
        doc,
        anchor,
        "实际结果与分析：",
        "同一 ACK 的聚合计数达到 12~17 次，同一数据序号被发送 5~7 次，证明 10% 丢包下出现了重复确认和重传，且传输没有在首次丢包后永久停滞。结合 duplicate_ack_count>=3 的实现可确认快速重传入口存在。不过截图中的 sort | uniq -c 会丢失原始时间顺序，无法仅凭聚合计数证明某次重发一定早于 RTO，因此该用例判为“丢包恢复通过、快速重传时机证据不完整”。复测时应保留连续的 RECV/SEND/RTTS 时间戳，并验证第三个重复 ACK 到重发之间的间隔小于当时的 TimeoutInterval。",
    )
    add_picture(doc, anchor, FIG_FAST, 6.0, "图表 8  10%丢包下重复累计ACK与相同序号重发统计")

    add_heading(doc, anchor, "6.6.3 RTT估计与RTO", level=3)
    add_labeled_paragraph(
        doc,
        anchor,
        "测试条件与流程：",
        "在基准链路上完成连接和数据传输，从 client.event.trace 提取 RTTS 事件，逐项核对 SampleRTT、EstimatedRTT、DeviationRTT 和 TimeoutInterval。",
    )
    add_labeled_paragraph(
        doc,
        anchor,
        "预期结果：",
        "首次采样令 SRTT=R、RTTVAR=R/2；后续采用 RTTVAR←0.75×RTTVAR+0.25×|SRTT-R|、SRTT←0.875×SRTT+0.125×R，并令 RTO=clamp(SRTT+4×RTTVAR, 200, 4000) ms。已重传报文不参与采样。",
    )
    add_labeled_paragraph(
        doc,
        anchor,
        "实际结果与分析：",
        "首个有效 SampleRTT 为 40.781982 ms，SRTT 同为 40.781982 ms，RTTVAR 为 20.390991 ms，原始 RTO 约 122.346 ms，因小于下界而记录为 200 ms。第二个样本为 14.782959 ms，计算得到 SRTT=37.532104 ms、RTTVAR=21.792999 ms，与 trace 完全一致。后续样本虽在约 15~106 ms 间波动，估计值仍平滑变化且 RTO 保持在 200 ms 下界。process_ack() 仅在报文未重传时取样，符合 Karn 算法。本项通过。",
    )
    add_picture(doc, anchor, FIG_RTT, 6.0, "图表 9  正常传输中的RTT采样、平滑估计与RTO下界")

    add_heading(doc, anchor, "6.6.4 小结", level=3)
    add_paragraph(
        doc,
        anchor,
        "现有证据验证了累计 ACK、失序后按序交付、丢包重发以及 RFC 6298 RTT/RTO 更新路径，但尚不能宣称可靠传输整体通过：文件逐字节比较失败，快速重传截图也缺少可用于区分 RTO 的原始时序。报告保留该失败结果，待完成首个差异字节定位和带时间戳复测后再更新验收结论。",
        first_indent=True,
    )


def build_section_67(doc, anchor):
    add_paragraph(
        doc,
        anchor,
        "流量控制测试关注接收端 16 位通告窗口、发送端窗口约束、接收缓存释放后的窗口更新，以及零窗口暂停与恢复。当前代码已实现相应控制路径，但所附运行证据主要针对可靠传输，没有记录窗口降至 0 的过程。因此本节同时给出可复现测试用例、已有静态/运行结果和未通过验收的原因。",
        first_indent=True,
    )

    add_result_table(
        doc,
        anchor,
        [
            (
                "RDT-FLOW-01\n通告窗口约束",
                "基准链路持续传输；同步提取服务端RWND与客户端SWND，并计算每时刻在途字节。",
                "rwnd等于可用缓存且不超过65535 B；发送右边界不超过snd_una+peer_rwnd。",
                "代码路径核验通过；现有截图未包含RWND/SWND时序及FlightSize，动态证据不足。",
            ),
            (
                "RDT-FLOW-02\n零窗口与恢复",
                "暂停服务端tju_recv，使超过接收容量的数据进入；观察rwnd=0和探测；随后恢复读取。",
                "零窗口期间停止发送新数据；周期发送探测；窗口更新后自动续传且文件一致。",
                "现有RDT接收端持续读取，未形成rwnd=0；未观察到零窗口探测。未完成动态验收。",
            ),
        ],
    )

    add_heading(doc, anchor, "6.7.1 通告窗口与发送约束", level=3)
    add_labeled_paragraph(
        doc,
        anchor,
        "测试条件与流程：",
        "在 100 Mbps、20 ms、0% 丢包的基准链路上持续发送数据，同时从 server.event.trace 提取 RWND、从 client.event.trace 提取 SWND；按 ACK 推进计算 snd_una，并核对每次发送后的最高序号是否超过 snd_una+peer_rwnd。接收端应用持续调用 tju_recv()，用于观察缓存占用和释放时窗口的下降与回升。",
    )
    add_labeled_paragraph(
        doc,
        anchor,
        "预期结果：",
        "通告窗口等于 recv_capacity 减去已按序缓存字节和失序缓存字节，并截断到 16 位最大值 65535 B；发送端只发送位于 snd_una+peer_rwnd 右边界以内的数据。应用读取释放缓存后，接收端应立即发送新的 ACK/窗口通告。",
    )
    add_labeled_paragraph(
        doc,
        anchor,
        "实际结果与分析：",
        "静态核验表明 advertised_window_locked() 按 received_len+recv_ooo_bytes 计算占用并将结果限制为 65535 B；flush_send_queue() 使用 snd_una+peer_rwnd 作为发送右边界；收到报文后 peer_rwnd 由首部 advertised_window 更新；tju_recv() 释放缓存、排入可连续交付的失序段后发送 ACK。因此设计和实现路径满足基本 rwnd 约束。可是现有截图没有 RWND/SWND 序列，也没有逐时刻 FlightSize，无法用运行数据证明约束始终成立。本项只能记为“静态核验通过、动态验收待补”。",
    )

    add_heading(doc, anchor, "6.7.2 零窗口、探测与恢复", level=3)
    add_labeled_paragraph(
        doc,
        anchor,
        "测试条件与流程：",
        "构造专用接收端：建立连接后暂停调用 tju_recv()，发送端传输量超过接收缓存容量 5000×1375 B，直至 trace 出现 RWND size:0；继续等待至少一个当前 RTO，检查发送端是否停止普通数据并发送 1 B 探测；随后恢复 tju_recv()，观察窗口更新、发送恢复及最终文件比较。为缩短测试时间，也可仅在测试构建中减小 recv_capacity，但不得修改协议逻辑。",
    )
    add_labeled_paragraph(
        doc,
        anchor,
        "预期结果：",
        "peer_rwnd=0 时发送队列不得继续发送新数据；timer_main() 按当前 RTO 周期发送探测段。接收端读取后通告非零窗口，发送端收到更新 ACK 后继续发送，所有字节最终按序且仅交付一次。RTO退避不得导致窗口恢复后长期停顿。",
    )
    add_labeled_paragraph(
        doc,
        anchor,
        "实际结果与分析：",
        "当前 RDT 服务端持续调用 tju_recv()，接收缓存被及时释放，所附日志中没有 RWND size:0，也没有能够与窗口恢复关联的 1 B 探测序列，因而未实际覆盖零窗口。代码中 peer_rwnd=0 会阻止普通数据发送，send_window_probe() 构造 snd_una-1 的 1 B 重复数据，timer_main() 按 RTO 周期触发探测；tju_recv() 在窗口扩大后主动发送 ACK。这些代码只能证明测试具备实现基础，不能代替端到端验证。本项记录为未完成，需按上述专用接收端流程补测，并保留 RWND、SWND、SEND、RECV、DELV 与最终 cmp 证据。",
    )

    add_heading(doc, anchor, "6.7.3 流量控制结论", level=3)
    add_paragraph(
        doc,
        anchor,
        "接收窗口计算、发送右边界和窗口释放通知已经落实到代码，但当前证据不足以确认动态窗口约束和零窗口恢复。尤其在 6.6 的文件一致性仍失败时，不能仅凭窗口代码路径宣称流量控制通过。后续验收应先修复可靠传输的数据一致性，再执行暂停读取—零窗口—探测—恢复的完整闭环，并把两端 trace 与文件哈希纳入证据。",
        first_indent=True,
    )


def main():
    for path in (INPUT, FIG_REORDER, FIG_RTT, FIG_FAST):
        if not path.exists():
            raise FileNotFoundError(path)

    doc = Document(INPUT)
    clear_section_body(doc, "6.6可靠传输测试", "6.7流量控制测试")
    clear_section_body(doc, "6.7流量控制测试", "6.8 代表性AI协作与验证记录")
    update_summary_table(doc)

    section_67 = find_paragraph(doc, "6.7流量控制测试")
    section_68 = find_paragraph(doc, "6.8 代表性AI协作与验证记录")
    build_section_66(doc, section_67)
    build_section_67(doc, section_68)

    doc.core_properties.title = "TJU_TCP课程报告（补全6.6可靠传输测试与6.7流量控制测试）"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
